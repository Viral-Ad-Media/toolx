from io import BytesIO
from xml.sax.saxutils import escape
import logging
from smtplib import SMTPException

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.paginator import Paginator
from django.core.mail import send_mail
from django.urls import reverse
from django.views.decorators.http import require_http_methods
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_decode, urlsafe_base64_encode
from docx import Document
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate

from .forms import InstantGeneratorForm, ParaphraseForm, ProfileForm, SignUpForm, UserForm, ActivationResendForm
from .models import InstantGenerator, Paraphrase, Profile
from .tokens import account_activation_token

User = get_user_model()


logger = logging.getLogger(__name__)


def send_activation(user):
    message = render_to_string('registration/activation_email.html', {
        'user': user,
        'activation_url': settings.PUBLIC_ORIGIN + reverse('activate', args=[
            urlsafe_base64_encode(force_bytes(user.pk)), account_activation_token.make_token(user),
        ]),
    })
    if send_mail('Activate Your Tool-X Account', message, settings.DEFAULT_FROM_EMAIL, [user.email]) != 1:
        raise SMTPException('Activation delivery was not accepted.')


@require_http_methods(['GET', 'POST'])
def signup(request):
    form = SignUpForm(request.POST if request.method == 'POST' else None)
    if request.method == 'POST' and form.is_valid():
        try:
            with transaction.atomic():
                user = form.save(commit=False)
                user.is_active = False
                user.save()
                Profile.objects.filter(user=user).update(activation_pending=True)
                # Avoid using the cached profile from the signal before updating state.
                user.refresh_from_db()
                send_activation(user)
        except (SMTPException, OSError):
            logger.warning('Registration activation email delivery failed.')
            form.add_error(None, 'We could not send your activation email. Please try again later.')
        else:
            return redirect('activation_sent')
    return render(request, 'registration/signup.html', {'form': form})


def activation_sent(request):
    return render(request, 'registration/activation_sent.html', {})


@require_http_methods(['GET', 'POST'])
def activation_resend(request):
    form = ActivationResendForm(request.POST if request.method == 'POST' else None)
    if request.method == 'POST' and form.is_valid():
        user = User.objects.filter(email__iexact=form.cleaned_data['email'], is_active=False,
                                   profile__activation_pending=True, profile__email_confirmed=False).first()
        if user:
            try:
                send_activation(user)
            except (SMTPException, OSError):
                logger.warning('Activation resend delivery failed.')
        # Same response for nonexistent, active, suspended, and pending identities.
        return redirect('activation_sent')
    return render(request, 'registration/activation_resend.html', {'form': form})


@require_http_methods(['GET'])
@transaction.atomic
def activate(request, uidb64, token):
    try:
        uid = urlsafe_base64_decode(uidb64).decode()
        user = User.objects.select_for_update().get(pk=uid)
    except (TypeError, ValueError, OverflowError, UnicodeDecodeError, User.DoesNotExist):
        user = None
    if user is None or user.is_active:
        return render(request, 'registration/activation_invalid.html', status=400)
    profile_obj = Profile.objects.select_for_update().filter(user=user).first()
    if not profile_obj or not profile_obj.activation_pending or profile_obj.email_confirmed or not account_activation_token.check_token(user, token):
        return render(request, 'registration/activation_invalid.html', status=400)
    profile_obj.activation_pending = False
    profile_obj.email_confirmed = True
    profile_obj.save(update_fields=['activation_pending', 'email_confirmed'])
    user.is_active = True
    user.save(update_fields=['is_active'])
    messages.success(request, 'Your email is confirmed. Please sign in.')
    return redirect('login')


@login_required
def profile(request):
    return render(request, 'registration/profile.html', {'user': request.user})


@login_required
@transaction.atomic
def edit_profile(request):
    profile_obj, _ = Profile.objects.get_or_create(user=request.user)

    if request.method == 'POST':
        user_form = UserForm(request.POST, instance=request.user)
        form = ProfileForm(request.POST, request.FILES, instance=profile_obj)
        if user_form.is_valid() and form.is_valid():
            user_form.save()
            form.save()
            messages.success(request, 'Your profile was successfully updated.')
            return redirect('profile')
        messages.error(request, 'Please correct the error below.')
    else:
        user_form = UserForm(instance=request.user)
        form = ProfileForm(instance=profile_obj)

    return render(
        request,
        'registration/edit_profile.html',
        {
            'user_form': user_form,
            'form': form,
        },
    )


@login_required
def dashboard(request):
    context = {
        'adcopies': InstantGenerator.objects.filter(user=request.user).only('Get_Attention', 'created_on').order_by('-created_on', '-pk')[:6],
        'my_paraphrase': Paraphrase.objects.filter(user=request.user).only('Title', 'created_on').order_by('-created_on', '-pk')[:6],
    }
    return render(request, 'instant_generator/dashboard.html', context)


@login_required
@transaction.atomic
def create(request):
    if request.method == 'POST':
        form = InstantGeneratorForm(request.POST)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.user = request.user
            instance.save()
            return redirect('congratulation')
    else:
        form = InstantGeneratorForm()

    return render(request, 'instant_generator/create.html', {'form': form})


@login_required
@transaction.atomic
def create_paraphrase(request):
    if request.method == 'POST':
        form = ParaphraseForm(request.POST)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.user = request.user
            instance.save()
            return redirect('congratulation')
    else:
        form = ParaphraseForm()

    return render(request, 'instant_generator/create_paraphrase.html', {'form': form})


@login_required
def congratulation(request):
    return render(request, 'instant_generator/congratulation.html', {})


@login_required
def paraphrase(request):
    my_paraphrase = Paginator(Paraphrase.objects.filter(user=request.user).only('Title', 'created_on').order_by('-created_on', '-pk'), 20).get_page(request.GET.get('page'))
    return render(request, 'instant_generator/paraphrase.html', {'my_paraphrase': my_paraphrase})


@login_required
def my_adcopies(request):
    adcopies = Paginator(InstantGenerator.objects.filter(user=request.user).only('Get_Attention', 'created_on').order_by('-created_on', '-pk'), 20).get_page(request.GET.get('page'))
    return render(request, 'instant_generator/my_adcopies.html', {'adcopies': adcopies})


@login_required
def preview(request, pk):
    generated = get_object_or_404(InstantGenerator, pk=pk, user=request.user)
    return render(request, 'instant_generator/preview.html', {'generated': generated})


@login_required
def paraphrase_preview(request, pk):
    original = get_object_or_404(Paraphrase, pk=pk, user=request.user)
    context = {
        'original': original,
    }
    return render(request, 'instant_generator/paraphrase_preview.html', context)


@login_required
def pdf(request, pk):
    generated = get_object_or_404(InstantGenerator, pk=pk, user=request.user)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Sales Letter.pdf"'

    if any(len(getattr(generated, f.name)) > 10000 for f in generated._meta.fields if f.get_internal_type() in ('CharField', 'TextField')):
        return HttpResponse('This record exceeds the export size limit. Shorten the content first.', status=400)

    pdf_buffer = BytesIO()
    my_doc = SimpleDocTemplate(pdf_buffer)
    sample_style_sheet = getSampleStyleSheet()

    sections = [
        (generated.Get_Attention, 'Title'),
        (generated.Identify_the_Problem_Your_Audience_Have, 'BodyText'),
        (generated.Provide_the_Solution, 'BodyText'),
        (generated.Present_your_Credentials, 'BodyText'),
        (generated.Show_the_Benefits, 'BodyText'),
        (generated.Give_Social_Proof, 'BodyText'),
        (generated.Make_Your_Offer, 'Heading4'),
        (generated.Give_a_Guarantee, 'Italic'),
        (generated.Inject_Scarcity, 'BodyText'),
        (generated.Call_to_action, 'Heading3'),
        (generated.Give_a_Warning, 'BodyText'),
        (generated.Close_with_a_Reminder, 'Heading4'),
    ]
    my_doc.build([Paragraph(escape(text).replace('\n', '<br/>'), sample_style_sheet[style]) for text, style in sections])

    response.write(pdf_buffer.getvalue())
    pdf_buffer.close()
    return response


@login_required
def docx(request, pk):
    generated = get_object_or_404(InstantGenerator, pk=pk, user=request.user)
    if any(len(getattr(generated, f.name)) > 10000 for f in generated._meta.fields if f.get_internal_type() in ('CharField', 'TextField')):
        return HttpResponse('This record exceeds the export size limit. Shorten the content first.', status=400)
    document = Document()
    document.add_heading(generated.Get_Attention, 0)
    document.add_paragraph(generated.Identify_the_Problem_Your_Audience_Have)
    document.add_paragraph(generated.Provide_the_Solution)
    document.add_paragraph(generated.Present_your_Credentials)
    document.add_paragraph(generated.Show_the_Benefits)
    document.add_paragraph(generated.Give_Social_Proof)
    document.add_paragraph(generated.Make_Your_Offer, style='Intense Quote')
    document.add_paragraph(generated.Give_a_Guarantee)
    document.add_paragraph(generated.Inject_Scarcity)
    document.add_paragraph(generated.Call_to_action)
    document.add_paragraph(generated.Give_a_Warning)
    document.add_paragraph().add_run(generated.Close_with_a_Reminder).bold = True

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="Sales Letter.docx"'
    document.save(response)
    return response



@login_required
@require_http_methods(['GET', 'POST'])
def edit_draft(request, pk):
    draft = get_object_or_404(Paraphrase, pk=pk, user=request.user)
    form = ParaphraseForm(request.POST if request.method == 'POST' else None, instance=draft)
    if request.method == 'POST' and form.is_valid():
        form.save()
        return redirect('paraphrase_preview', pk=draft.pk)
    return render(request, 'instant_generator/create_paraphrase.html', {'form': form})
